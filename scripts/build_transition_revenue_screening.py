#!/usr/bin/env python3
"""Build the SCR transition-risk annual-revenue screening trial.

The builder intentionally keeps three layers distinct:

1. SCR source impacts and ratings;
2. InfraSure-derived percentage-point factors; and
3. revenue snapshots calculated from the accepted V6 revenue layers.

The output Parquet is a research artifact. It is not a production cash-flow
forecast and it is never applied to TIV.
"""

from __future__ import annotations

import argparse
import concurrent.futures
import hashlib
import json
import math
import re
import shutil
import subprocess
import tempfile
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook


SOLAR_PREFIX = "gs://infrasure-scr-data/Solar-asset/transition_risk/"
WIND_PREFIX = "gs://infrasure-scr-data/onshore_wind/transition_risk/"
CELL_RE = re.compile(r"_Cell_(\d+)_")
WIND_REGIME_B_MIN_CELL = 267458
SCENARIOS = (
    "Below 2°C",
    "Climate Breakdown",
    "Climate Destabilization",
    "Current Policies",
    "Delayed transition",
    "Expected",
    "Fragmented World",
    "Nationally Determined Contributions (NDCs)",
    "Net Zero 2050",
)
HORIZONS = (2025, 2030, 2035, 2040, 2045, 2050)
INDICATOR_HORIZON_SEMANTICS = (
    "SCR indicator values are 10-year averages centered on the selected horizon."
)
FINANCIAL_IMPACT_HORIZON_SEMANTICS = (
    "SCR subrisk financial impacts are average annual impacts from 2025 through "
    "the selected horizon; they are neither cumulative dollar totals nor "
    "point-in-time-only impacts."
)
SIGN_CONVENTION = (
    "Positive adjusted impact is treated as an annual-revenue uplift and negative "
    "adjusted impact as an annual-revenue drag."
)
REVENUE_KWP = "annual_revenue_adj_curtailment_market_value_p75_usd_kwp_yr"
REVENUE_100MW = (
    "annual_revenue_adj_curtailment_market_value_p75_100mw_usd_yr"
)
REVENUE_COLUMNS = (
    "cell_id",
    "lat_center",
    "lon_center",
    "state_abbr",
    REVENUE_KWP,
    REVENUE_100MW,
)
NUMERIC_OUTPUT_COLUMNS = (
    "carbon_price",
    "scope_1_2_intensity",
    "revenue_growth",
    "scope_3_intensity",
    "direct_carbon_cost_impact_raw",
    "direct_carbon_cost_impact_adjusted",
    "market_demand_impact_raw",
    "market_demand_impact_adjusted",
    "baseline_revenue_usd_kwp_yr",
    "baseline_revenue_100mw_usd_yr",
    "direct_carbon_fraction",
    "direct_carbon_factor",
    "direct_carbon_adjusted_revenue_usd_kwp_yr",
    "direct_carbon_revenue_delta_usd_kwp_yr",
    "direct_carbon_adjusted_revenue_100mw_usd_yr",
    "direct_carbon_revenue_delta_100mw_usd_yr",
    "market_demand_fraction",
    "market_demand_factor",
    "market_demand_adjusted_revenue_usd_kwp_yr",
    "market_demand_revenue_delta_usd_kwp_yr",
    "market_demand_adjusted_revenue_100mw_usd_yr",
    "market_demand_revenue_delta_100mw_usd_yr",
    "overall_transition_impact",
    "overall_transition_fraction",
    "overall_transition_screening_factor",
    "overall_adjusted_revenue_usd_kwp_yr",
    "overall_revenue_delta_usd_kwp_yr",
    "overall_adjusted_revenue_100mw_usd_yr",
    "overall_revenue_delta_100mw_usd_yr",
)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def sha256_lines(lines: Iterable[str]) -> str:
    normalized = "\n".join(sorted(line.strip() for line in lines if line.strip()))
    return hashlib.sha256((normalized + "\n").encode()).hexdigest()


def list_gcs(prefix: str) -> list[str]:
    result = subprocess.run(
        ["gcloud", "storage", "ls", prefix],
        check=True,
        capture_output=True,
        text=True,
    )
    return [line.strip() for line in result.stdout.splitlines() if line.strip()]


def cell_id_from_uri(uri: str) -> int:
    match = CELL_RE.search(uri)
    if not match:
        raise ValueError(f"No cell ID in SCR object name: {uri}")
    return int(match.group(1))


def copy_gcs(uri: str, destination: Path) -> Path:
    destination.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ["gcloud", "storage", "cp", uri, str(destination)],
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    return destination


def evenly_spaced(items: list[str], count: int) -> list[str]:
    if count >= len(items):
        return list(items)
    indexes = {round(index * (len(items) - 1) / (count - 1)) for index in range(count)}
    return [items[index] for index in sorted(indexes)]


def normalized_workbook_hash(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = list(workbook["Output"].iter_rows(values_only=True))
    header = list(rows[0])
    keep = (
        "scenario",
        "timeHorizon",
        "indicator",
        "indicatorUnit",
        "indicatorValue",
        "subrisk",
        "subriskRevenueImpact",
        "adjustedSubriskRevenueImpact",
        "subriskExposureRating",
        "adjustedSubriskExposureRating",
        "transitionExposureRating",
        "adjustedTransitionExposureRating",
    )
    indexes = [header.index(column) for column in keep]
    normalized = [[row[index] for index in indexes] for row in rows[1:]]
    payload = json.dumps(normalized, default=str, separators=(",", ":"))
    return hashlib.sha256(payload.encode()).hexdigest()


def audit_templates(
    solar_uris: list[str],
    wind_uris: list[str],
    root: Path,
    sample_count: int,
) -> dict[str, Any]:
    solar_sample = evenly_spaced(solar_uris, sample_count)
    wind_sample = evenly_spaced(wind_uris, sample_count)
    for boundary in (267457, 267458):
        uri = next((item for item in wind_uris if f"_Cell_{boundary}_" in item), None)
        if uri and uri not in wind_sample:
            wind_sample.append(uri)

    downloads: list[tuple[str, Path, str]] = []
    for asset, uris in (("solar", solar_sample), ("wind", wind_sample)):
        for uri in uris:
            downloads.append((uri, root / f"{asset}_{Path(uri).name}", asset))

    def download(item: tuple[str, Path, str]) -> tuple[Path, str]:
        uri, destination, asset = item
        return copy_gcs(uri, destination), asset

    with concurrent.futures.ThreadPoolExecutor(max_workers=8) as executor:
        completed = list(executor.map(download, downloads))

    groups: dict[str, dict[str, list[int]]] = {
        "solar": defaultdict(list),
        "wind": defaultdict(list),
    }
    for path, asset in completed:
        groups[asset][normalized_workbook_hash(path)].append(cell_id_from_uri(path.name))

    if len(groups["solar"]) != 1:
        raise AssertionError(f"Solar template audit found {len(groups['solar'])} templates")
    if len(groups["wind"]) != 2:
        raise AssertionError(f"Wind template audit found {len(groups['wind'])} templates")
    for ids in groups["wind"].values():
        sides = {cell_id >= WIND_REGIME_B_MIN_CELL for cell_id in ids}
        if len(sides) != 1:
            raise AssertionError("A sampled Wind template crosses the regime boundary")

    return {
        asset: {
            "sampled_workbook_count": sum(len(ids) for ids in hashes.values()),
            "normalized_template_count": len(hashes),
            "templates": [
                {
                    "normalized_sha256": digest,
                    "sample_count": len(ids),
                    "minimum_cell_id": min(ids),
                    "maximum_cell_id": max(ids),
                }
                for digest, ids in sorted(hashes.items())
            ],
        }
        for asset, hashes in groups.items()
    }


def parse_template(path: Path) -> tuple[dict[tuple[str, int], dict[str, Any]], dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = list(workbook["Output"].iter_rows(values_only=True))
    header = list(rows[0])
    records = [dict(zip(header, row)) for row in rows[1:]]
    grouped: dict[tuple[str, int], list[dict[str, Any]]] = defaultdict(list)
    for row in records:
        grouped[(str(row["scenario"]), int(row["timeHorizon"]))].append(row)

    parsed: dict[tuple[str, int], dict[str, Any]] = {}
    for key, values in grouped.items():
        by_indicator = {str(row["indicator"]): row for row in values}
        direct = by_indicator["Carbon price"]
        market = by_indicator["Revenue growth"]
        parsed[key] = {
            "carbon_price": direct["indicatorValue"],
            "carbon_price_unit": direct["indicatorUnit"],
            "scope_1_2_intensity": by_indicator["Scope 1&2 emissions intensity"]["indicatorValue"],
            "scope_1_2_intensity_unit": by_indicator["Scope 1&2 emissions intensity"]["indicatorUnit"],
            "revenue_growth": market["indicatorValue"],
            "revenue_growth_unit": market["indicatorUnit"],
            "scope_3_intensity": by_indicator["Scope 3 emissions intensity"]["indicatorValue"],
            "scope_3_intensity_unit": by_indicator["Scope 3 emissions intensity"]["indicatorUnit"],
            "direct_carbon_cost_impact_raw": direct["subriskRevenueImpact"],
            "direct_carbon_cost_impact_adjusted": direct["adjustedSubriskRevenueImpact"],
            "direct_carbon_cost_rating_raw": direct["subriskExposureRating"],
            "direct_carbon_cost_rating_adjusted": direct["adjustedSubriskExposureRating"],
            "market_demand_impact_raw": market["subriskRevenueImpact"],
            "market_demand_impact_adjusted": market["adjustedSubriskRevenueImpact"],
            "market_demand_rating_raw": market["subriskExposureRating"],
            "market_demand_rating_adjusted": market["adjustedSubriskExposureRating"],
            "overall_transition_rating_raw": direct["transitionExposureRating"],
            "overall_transition_rating_adjusted": direct["adjustedTransitionExposureRating"],
        }

    observed_scenarios = tuple(dict.fromkeys(key[0] for key in grouped))
    observed_horizons = tuple(sorted({key[1] for key in grouped}))
    if set(observed_scenarios) != set(SCENARIOS):
        raise AssertionError(f"Unexpected scenarios in {path}: {observed_scenarios}")
    if observed_horizons != HORIZONS:
        raise AssertionError(f"Unexpected horizons in {path}: {observed_horizons}")

    first = records[0]
    metadata = {
        "source_uri": None,
        "source_file_sha256": sha256_file(path),
        "normalized_template_sha256": normalized_workbook_hash(path),
        "report_date": first["reportDate"].date().isoformat(),
        "ticcs_subclass": first["ticcsSubClass"],
        "ticcs_subclass_name": first["ticcsSubClassName"],
        "row_count": len(records),
    }
    return parsed, metadata


def nullable_float(value: Any) -> float | None:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    number = float(value)
    return number if math.isfinite(number) else None


def factor(impact: float | None) -> float | None:
    return None if impact is None else 1.0 + impact / 100.0


def product(value: float | None, multiplier: float | None) -> float | None:
    return None if value is None or multiplier is None else value * multiplier


def difference(value: float | None, base: float | None) -> float | None:
    return None if value is None or base is None else value - base


def load_revenue_layers(solar_path: Path, wind_path: Path) -> dict[str, pd.DataFrame]:
    solar = pd.read_parquet(solar_path, columns=list(REVENUE_COLUMNS))
    wind = pd.read_parquet(wind_path, columns=list(REVENUE_COLUMNS))
    if solar["cell_id"].duplicated().any() or wind["cell_id"].duplicated().any():
        raise AssertionError("Revenue layer contains duplicate cell IDs")
    if len(solar) != 13085 or len(wind) != 13083:
        raise AssertionError(f"Unexpected V6 revenue counts: Solar={len(solar)}, Wind={len(wind)}")

    canonical = solar[["cell_id", "lat_center", "lon_center", "state_abbr"]].copy()
    wind_values = wind[["cell_id", REVENUE_KWP, REVENUE_100MW]].copy()
    wind_complete = canonical.merge(wind_values, on="cell_id", how="left", validate="one_to_one")
    return {"solar": solar, "wind_farm": wind_complete}


def regime_for(asset: str, cell_id: int, scr_ids: set[int]) -> str | None:
    if cell_id not in scr_ids:
        return None
    if asset == "solar":
        return "solar_v1"
    return "wind_a" if cell_id < WIND_REGIME_B_MIN_CELL else "wind_b"


def availability_status(scr_available: bool, revenue_available: bool) -> str:
    if not scr_available:
        return "missing_scr"
    if not revenue_available:
        return "missing_revenue"
    return "available"


def build_batch(
    asset: str,
    revenue: pd.DataFrame,
    scr_ids: set[int],
    templates: dict[str, dict[tuple[str, int], dict[str, Any]]],
    template_meta: dict[str, dict[str, Any]],
    scenario: str,
    horizon: int,
    revenue_sha: str,
    run_id: str,
) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for source in revenue.itertuples(index=False):
        cell_id = int(source.cell_id)
        regime = regime_for(asset, cell_id, scr_ids)
        scr_available = regime is not None
        baseline_kwp = nullable_float(getattr(source, REVENUE_KWP))
        baseline_100mw = nullable_float(getattr(source, REVENUE_100MW))
        revenue_available = baseline_kwp is not None and baseline_100mw is not None
        template = templates[regime][(scenario, horizon)] if regime else None
        direct_raw = nullable_float(template["direct_carbon_cost_impact_raw"]) if template else None
        direct_adjusted = nullable_float(template["direct_carbon_cost_impact_adjusted"]) if template else None
        market_raw = nullable_float(template["market_demand_impact_raw"]) if template else None
        market_adjusted = nullable_float(template["market_demand_impact_adjusted"]) if template else None
        overall = (
            direct_adjusted + market_adjusted
            if direct_adjusted is not None and market_adjusted is not None
            else None
        )
        direct_factor = factor(direct_adjusted)
        market_factor = factor(market_adjusted)
        overall_factor = factor(overall)
        direct_kwp = product(baseline_kwp, direct_factor)
        market_kwp = product(baseline_kwp, market_factor)
        overall_kwp = product(baseline_kwp, overall_factor)
        direct_100mw = product(baseline_100mw, direct_factor)
        market_100mw = product(baseline_100mw, market_factor)
        overall_100mw = product(baseline_100mw, overall_factor)
        rows.append(
            {
                "cell_id": cell_id,
                "lat_center": float(source.lat_center),
                "lon_center": float(source.lon_center),
                "state_abbr": None if source.state_abbr is None else str(source.state_abbr),
                "asset_type": asset,
                "ticcs_subclass": template_meta[regime]["ticcs_subclass"] if regime else None,
                "scenario": scenario,
                "horizon": horizon,
                "scr_report_date": template_meta[regime]["report_date"] if regime else None,
                "scr_source_regime": regime,
                "scr_available": scr_available,
                "revenue_available": revenue_available,
                "availability_status": availability_status(scr_available, revenue_available),
                "carbon_price": template["carbon_price"] if template else None,
                "carbon_price_unit": template["carbon_price_unit"] if template else None,
                "scope_1_2_intensity": template["scope_1_2_intensity"] if template else None,
                "scope_1_2_intensity_unit": template["scope_1_2_intensity_unit"] if template else None,
                "revenue_growth": template["revenue_growth"] if template else None,
                "revenue_growth_unit": template["revenue_growth_unit"] if template else None,
                "scope_3_intensity": template["scope_3_intensity"] if template else None,
                "scope_3_intensity_unit": template["scope_3_intensity_unit"] if template else None,
                "direct_carbon_cost_impact_raw": direct_raw,
                "direct_carbon_cost_impact_adjusted": direct_adjusted,
                "direct_carbon_cost_rating_raw": template["direct_carbon_cost_rating_raw"] if template else None,
                "direct_carbon_cost_rating_adjusted": template["direct_carbon_cost_rating_adjusted"] if template else None,
                "market_demand_impact_raw": market_raw,
                "market_demand_impact_adjusted": market_adjusted,
                "market_demand_rating_raw": template["market_demand_rating_raw"] if template else None,
                "market_demand_rating_adjusted": template["market_demand_rating_adjusted"] if template else None,
                "overall_transition_rating_raw": template["overall_transition_rating_raw"] if template else None,
                "overall_transition_rating_adjusted": template["overall_transition_rating_adjusted"] if template else None,
                "baseline_revenue_usd_kwp_yr": baseline_kwp,
                "baseline_revenue_100mw_usd_yr": baseline_100mw,
                "direct_carbon_fraction": None if direct_adjusted is None else direct_adjusted / 100.0,
                "direct_carbon_factor": direct_factor,
                "direct_carbon_adjusted_revenue_usd_kwp_yr": direct_kwp,
                "direct_carbon_revenue_delta_usd_kwp_yr": difference(direct_kwp, baseline_kwp),
                "direct_carbon_adjusted_revenue_100mw_usd_yr": direct_100mw,
                "direct_carbon_revenue_delta_100mw_usd_yr": difference(direct_100mw, baseline_100mw),
                "market_demand_fraction": None if market_adjusted is None else market_adjusted / 100.0,
                "market_demand_factor": market_factor,
                "market_demand_adjusted_revenue_usd_kwp_yr": market_kwp,
                "market_demand_revenue_delta_usd_kwp_yr": difference(market_kwp, baseline_kwp),
                "market_demand_adjusted_revenue_100mw_usd_yr": market_100mw,
                "market_demand_revenue_delta_100mw_usd_yr": difference(market_100mw, baseline_100mw),
                "overall_transition_impact": overall,
                "overall_transition_fraction": None if overall is None else overall / 100.0,
                "overall_transition_screening_factor": overall_factor,
                "overall_adjusted_revenue_usd_kwp_yr": overall_kwp,
                "overall_revenue_delta_usd_kwp_yr": difference(overall_kwp, baseline_kwp),
                "overall_adjusted_revenue_100mw_usd_yr": overall_100mw,
                "overall_revenue_delta_100mw_usd_yr": difference(overall_100mw, baseline_100mw),
                "impact_scale_assumption": "source_value_is_percentage_points",
                "direct_carbon_application_status": "experimental_annual_revenue_proxy",
                "overall_combination_method": "additive_same_revenue_base_proxy",
                "tiv_application_status": "not_applied",
                "revenue_release_id": "index_analysis_lab_v6_stable_delivery_alias",
                "revenue_source_sha256": revenue_sha,
                "scr_trial_run_id": run_id,
            }
        )
    frame = pd.DataFrame.from_records(rows)
    # Some 2025 scenario batches contain structural nulls for an entire impact
    # column. Force a stable floating-point contract so the first Parquet row
    # group cannot lock those fields to Arrow's null-only type.
    for column in NUMERIC_OUTPUT_COLUMNS:
        frame[column] = pd.to_numeric(frame[column], errors="coerce").astype("float64")
    return frame


def json_value(value: Any) -> Any:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    if hasattr(value, "item"):
        return value.item()
    return value


def write_dashboard_bundle(
    output_dir: Path,
    revenue_layers: dict[str, pd.DataFrame],
    scr_ids: dict[str, set[int]],
    templates: dict[str, dict[tuple[str, int], dict[str, Any]]],
    template_meta: dict[str, dict[str, Any]],
    run_id: str,
    generated_at: str,
    revenue_shas: dict[str, str],
    inventory_shas: dict[str, str],
) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    for asset, revenue in revenue_layers.items():
        rows = []
        for source in revenue.itertuples(index=False):
            cell_id = int(source.cell_id)
            regime = regime_for(asset, cell_id, scr_ids[asset])
            baseline_kwp = nullable_float(getattr(source, REVENUE_KWP))
            baseline_100mw = nullable_float(getattr(source, REVENUE_100MW))
            rows.append(
                {
                    "cell_id": cell_id,
                    "lat_center": float(source.lat_center),
                    "lon_center": float(source.lon_center),
                    "state_abbr": None if source.state_abbr is None else str(source.state_abbr),
                    "baseline_revenue_usd_kwp_yr": baseline_kwp,
                    "baseline_revenue_100mw_usd_yr": baseline_100mw,
                    "scr_source_regime": regime,
                    "scr_available": regime is not None,
                    "revenue_available": baseline_kwp is not None and baseline_100mw is not None,
                }
            )
        payload = {
            "schema_version": "scr-transition-screening-grid/v1",
            "asset": asset,
            "run_id": run_id,
            "rows": rows,
        }
        (output_dir / f"delivery_transition.grid.{asset}.json").write_text(
            json.dumps(payload, separators=(",", ":")), encoding="utf-8"
        )

    template_payload = {
        "schema_version": "scr-transition-screening-templates/v1",
        "run_id": run_id,
        "templates": {
            regime: {
                "metadata": template_meta[regime],
                "records": [
                    {"scenario": scenario, "horizon": horizon, **{key: json_value(value) for key, value in record.items()}}
                    for (scenario, horizon), record in values.items()
                ],
            }
            for regime, values in templates.items()
        },
    }
    (output_dir / "delivery_transition.templates.json").write_text(
        json.dumps(template_payload, separators=(",", ":")), encoding="utf-8"
    )

    manifest = {
        "schema_version": "scr-transition-screening-manifest/v1",
        "run_id": run_id,
        "generated_at": generated_at,
        "status": "research_trial",
        "claim_grade": "experimental_annual_revenue_proxy_not_production_cash_flow",
        "scenarios": list(SCENARIOS),
        "horizons": list(HORIZONS),
        "drivers": ["direct_carbon", "market_demand", "overall"],
        "default": {"asset": "wind_farm", "scenario": "Expected", "horizon": 2030, "driver": "overall", "metric": "adjusted_revenue"},
        "factor_definition": "1 + adjusted impact / 100",
        "overall_combination_method": "additive_same_revenue_base_proxy",
        "impact_scale_assumption": "source_value_is_percentage_points",
        "horizon_semantics": {
            "indicators": INDICATOR_HORIZON_SEMANTICS,
            "financial_impacts": FINANCIAL_IMPACT_HORIZON_SEMANTICS,
        },
        "sign_convention": SIGN_CONVENTION,
        "adjustment_policy": (
            "Use SCR adjusted subrisk impacts operationally and retain raw impacts "
            "for audit. Raw and adjusted impacts are identical in the current sources."
        ),
        "outlier_policy": "No clipping, compression, winsorization, or other outlier transform is applied.",
        "tiv_application_status": "not_applied",
        "revenue_release_id": "index_analysis_lab_v6_stable_delivery_alias",
        "assets": [
            {
                "asset_id": asset,
                "asset_label": "Solar" if asset == "solar" else "Wind Farm",
                "cell_count": len(revenue),
                "scr_available_cell_count": len(scr_ids[asset]),
                "revenue_available_cell_count": int(revenue[REVENUE_KWP].notna().sum()),
                "revenue_source_sha256": revenue_shas[asset],
                "scr_inventory_sha256": inventory_shas[asset],
                "grid_file": f"delivery_transition.grid.{asset}.json",
            }
            for asset, revenue in revenue_layers.items()
        ],
        "template_file": "delivery_transition.templates.json",
        "limitations": [
            "Research screening proxy; not a production cash-flow forecast.",
            "Direct Carbon Cost is temporarily applied to annual revenue pending denominator validation.",
            "SCR impact values are interpreted as percentage points for this trial.",
            "Missing SCR or revenue values remain unavailable and are never replaced with zero or 1.0x.",
            "The InfraSure numeric overall factor is not SCR's A-G overall transition rating.",
            FINANCIAL_IMPACT_HORIZON_SEMANTICS,
        ],
    }
    (output_dir / "delivery.transition.manifest.json").write_text(
        json.dumps(manifest, indent=2) + "\n", encoding="utf-8"
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    repo = Path(__file__).resolve().parents[1]
    parser.add_argument("--solar-revenue", type=Path, required=True)
    parser.add_argument("--wind-revenue", type=Path, required=True)
    parser.add_argument("--output-root", type=Path, default=repo / "runs" / "transition_revenue_screening")
    parser.add_argument("--dashboard-data-dir", type=Path)
    parser.add_argument("--run-id")
    parser.add_argument("--audit-samples", type=int, default=31)
    args = parser.parse_args()

    run_id = args.run_id or datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    generated_at = datetime.now(timezone.utc).isoformat()
    run_dir = args.output_root / f"run_id={run_id}"
    run_dir.mkdir(parents=True, exist_ok=False)

    solar_uris = list_gcs(SOLAR_PREFIX)
    wind_uris = list_gcs(WIND_PREFIX)
    solar_ids = {cell_id_from_uri(uri) for uri in solar_uris}
    wind_ids = {cell_id_from_uri(uri) for uri in wind_uris}
    if len(solar_uris) != 13041 or len(solar_ids) != 13041:
        raise AssertionError("Solar transition inventory is not the expected 13,041 unique cells")
    if len(wind_uris) != 13085 or len(wind_ids) != 13085:
        raise AssertionError("Wind transition inventory is not the expected 13,085 unique cells")
    if sum(cell_id < WIND_REGIME_B_MIN_CELL for cell_id in wind_ids) != 3271:
        raise AssertionError("Wind Template A inventory count changed")
    if sum(cell_id >= WIND_REGIME_B_MIN_CELL for cell_id in wind_ids) != 9814:
        raise AssertionError("Wind Template B inventory count changed")

    with tempfile.TemporaryDirectory(prefix="scr-transition-build-") as temp_name:
        temp = Path(temp_name)
        audit = audit_templates(solar_uris, wind_uris, temp / "audit", args.audit_samples)
        selected = {
            "solar_v1": solar_uris[0],
            "wind_a": next(uri for uri in wind_uris if cell_id_from_uri(uri) < WIND_REGIME_B_MIN_CELL),
            "wind_b": next(uri for uri in wind_uris if cell_id_from_uri(uri) == WIND_REGIME_B_MIN_CELL),
        }
        templates: dict[str, dict[tuple[str, int], dict[str, Any]]] = {}
        template_meta: dict[str, dict[str, Any]] = {}
        for regime, uri in selected.items():
            local = copy_gcs(uri, temp / "templates" / f"{regime}.xlsx")
            templates[regime], template_meta[regime] = parse_template(local)
            template_meta[regime]["source_uri"] = uri

    if template_meta["solar_v1"]["normalized_template_sha256"] not in {
        item["normalized_sha256"] for item in audit["solar"]["templates"]
    }:
        raise AssertionError("Selected Solar template did not match sampled Solar template")
    for regime in ("wind_a", "wind_b"):
        if template_meta[regime]["normalized_template_sha256"] not in {
            item["normalized_sha256"] for item in audit["wind"]["templates"]
        }:
            raise AssertionError(f"Selected {regime} template did not match sampled Wind templates")

    revenue_layers = load_revenue_layers(args.solar_revenue, args.wind_revenue)
    revenue_shas = {
        "solar": sha256_file(args.solar_revenue),
        "wind_farm": sha256_file(args.wind_revenue),
    }
    inventory_shas = {
        "solar": sha256_lines(solar_uris),
        "wind_farm": sha256_lines(wind_uris),
    }
    scr_ids = {"solar": solar_ids, "wind_farm": wind_ids}

    parquet_path = run_dir / "scr_transition_revenue_screening_conus.parquet"
    writer: pq.ParquetWriter | None = None
    row_count = 0
    status_counts: dict[str, int] = defaultdict(int)
    raw_adjusted_mismatch = 0
    try:
        for asset, revenue in revenue_layers.items():
            for scenario in SCENARIOS:
                for horizon in HORIZONS:
                    batch = build_batch(
                        asset,
                        revenue,
                        scr_ids[asset],
                        templates,
                        template_meta,
                        scenario,
                        horizon,
                        revenue_shas[asset],
                        run_id,
                    )
                    row_count += len(batch)
                    for key, count in batch["availability_status"].value_counts().items():
                        status_counts[str(key)] += int(count)
                    raw_adjusted_mismatch += int(
                        (
                            ~batch["direct_carbon_cost_impact_raw"].fillna(float("inf")).eq(
                                batch["direct_carbon_cost_impact_adjusted"].fillna(float("inf"))
                            )
                        ).sum()
                    )
                    raw_adjusted_mismatch += int(
                        (
                            ~batch["market_demand_impact_raw"].fillna(float("inf")).eq(
                                batch["market_demand_impact_adjusted"].fillna(float("inf"))
                            )
                        ).sum()
                    )
                    table = pa.Table.from_pandas(batch, preserve_index=False)
                    if writer is None:
                        writer = pq.ParquetWriter(parquet_path, table.schema, compression="zstd")
                    writer.write_table(table)
    finally:
        if writer is not None:
            writer.close()

    expected_rows = 2 * 13085 * len(SCENARIOS) * len(HORIZONS)
    if row_count != expected_rows:
        raise AssertionError(f"Expected {expected_rows:,} rows, wrote {row_count:,}")
    if raw_adjusted_mismatch:
        raise AssertionError(f"Found {raw_adjusted_mismatch} raw/adjusted mismatches")

    qa = {
        "schema_version": "scr-transition-revenue-screening-qa/v1",
        "run_id": run_id,
        "generated_at": generated_at,
        "row_count": row_count,
        "unique_grain_count": row_count,
        "asset_cell_count": {asset: len(frame) for asset, frame in revenue_layers.items()},
        "scr_inventory_cell_count": {"solar": len(solar_ids), "wind_farm": len(wind_ids)},
        "availability_row_counts": dict(sorted(status_counts.items())),
        "raw_adjusted_impact_mismatch_count": raw_adjusted_mismatch,
        "template_audit": audit,
        "wind_regime_counts": {
            "wind_a": sum(cell_id < WIND_REGIME_B_MIN_CELL for cell_id in wind_ids),
            "wind_b": sum(cell_id >= WIND_REGIME_B_MIN_CELL for cell_id in wind_ids),
        },
        "parquet_sha256": sha256_file(parquet_path),
        "checks": {
            "expected_rows": True,
            "unique_inventory_ids": True,
            "expected_template_count": True,
            "raw_equals_adjusted_in_current_sources": True,
            "tiv_not_used": True,
            "no_factor_compression_or_outlier_transform": True,
        },
    }
    (run_dir / "qa_report.json").write_text(json.dumps(qa, indent=2) + "\n", encoding="utf-8")

    readme = f"""# SCR transition revenue screening trial

Run ID: `{run_id}`
Status: research source run; not a production cash-flow forecast

## Coverage

- Rows: `{row_count:,}`
- Grain: `cell_id x asset_type x scenario x horizon`
- Assets: Solar and Wind Farm
- Scenarios: `9`
- Horizons: `6` (`2025`, `2030`, `2035`, `2040`, `2045`, `2050`)

## Operational interpretation

- SCR operational field: `adjustedSubriskRevenueImpact`; the raw field is retained for audit.
- Impact scale: interpreted as percentage points.
- Sign: positive is an annual-revenue uplift; negative is an annual-revenue drag.
- Direct Carbon Cost: temporarily applied to annual revenue as a disclosed screening proxy pending a governed OpEx denominator.
- Market Demand Shift: applied to annual revenue.
- Overall screen: `1 + (direct_carbon_adjusted + market_demand_adjusted) / 100`, with both impacts applied additively to the same revenue base.
- TIV: never used.
- Outliers: no clipping, compression, winsorization, or other transform is applied.
- Missing SCR or revenue values: retained as unavailable, never replaced with zero or `1.0x`.

## Horizon meaning

- Indicator values are 10-year averages centered on the selected horizon.
- Financial impacts are average annual impacts from 2025 through the selected horizon.
- A horizon value is not a cumulative dollar total and is not a point-in-time-only impact.

The numeric overall factor is an InfraSure screening calculation, not SCR's A-G overall transition rating.
See `manifest.json` for pinned inputs and assumptions and `qa_report.json` for validation results.
"""
    (run_dir / "README.md").write_text(readme, encoding="utf-8")

    manifest = {
        "schema_version": "scr-transition-revenue-screening-run/v1",
        "run_id": run_id,
        "generated_at": generated_at,
        "status": "research_trial_complete",
        "parquet": {
            "filename": parquet_path.name,
            "sha256": qa["parquet_sha256"],
            "row_count": row_count,
            "grain": "cell_id x asset_type x scenario x horizon",
        },
        "qa_report_sha256": sha256_file(run_dir / "qa_report.json"),
        "readme_sha256": sha256_file(run_dir / "README.md"),
        "source_inventories": {
            "solar": {"prefix": SOLAR_PREFIX, "count": len(solar_ids), "sha256": inventory_shas["solar"]},
            "wind_farm": {"prefix": WIND_PREFIX, "count": len(wind_ids), "sha256": inventory_shas["wind_farm"]},
        },
        "revenue_sources": {
            "solar": {"path": str(args.solar_revenue), "sha256": revenue_shas["solar"]},
            "wind_farm": {"path": str(args.wind_revenue), "sha256": revenue_shas["wind_farm"]},
        },
        "template_sources": template_meta,
        "assumptions": {
            "impact_scale": "source_value_is_percentage_points",
            "direct_carbon_application": "experimental_annual_revenue_proxy",
            "market_demand_application": "annual_revenue",
            "overall_combination": "additive_same_revenue_base_proxy",
            "factor_formula": "1 + adjusted impact / 100",
            "horizon_semantics": {
                "indicators": INDICATOR_HORIZON_SEMANTICS,
                "financial_impacts": FINANCIAL_IMPACT_HORIZON_SEMANTICS,
            },
            "sign_convention": SIGN_CONVENTION,
            "adjustment_policy": "adjusted_operational_raw_retained_for_audit",
            "outlier_policy": "none",
            "missing_value_policy": "unavailable_not_zero_or_one",
            "tiv_application": "not_applied",
        },
    }
    (run_dir / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")

    if args.dashboard_data_dir:
        write_dashboard_bundle(
            args.dashboard_data_dir,
            revenue_layers,
            scr_ids,
            templates,
            template_meta,
            run_id,
            generated_at,
            revenue_shas,
            inventory_shas,
        )

    print(json.dumps({"run_dir": str(run_dir), "parquet": str(parquet_path), "qa": qa}, indent=2))


if __name__ == "__main__":
    main()
