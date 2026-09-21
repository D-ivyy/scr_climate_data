#!/usr/bin/env python3
"""Inventory SCR GCS exports and profile deterministic workbook samples.

This script performs no GCS writes. It is deliberately small and auditable so
the four current source surfaces can be rechecked before production processing.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import subprocess
import tempfile
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DIR = Path(__file__).resolve().parent / "outputs"
CANONICAL_GRID_URI = (
    "gs://infrasure-benchmark/hazard_conus_grid/dev/common/benchmark_grid/"
    "served_conus_cell_ids_v2026_06.csv"
)
SURFACES = {
    "solar_physical": "gs://infrasure-scr-data/Solar-asset/physical_risk/",
    "solar_transition": "gs://infrasure-scr-data/Solar-asset/transition_risk/",
    "wind_physical": "gs://infrasure-scr-data/onshore_wind/physical_risk/",
    "wind_transition": "gs://infrasure-scr-data/onshore_wind/transition_risk/",
}
CELL_PATTERN = re.compile(r"CONUS13K_Cell_(\d+)_")
SAMPLE_QUANTILES = (0.0, 0.1, 0.25, 0.5, 0.75, 0.9, 1.0)


def run_text(args: list[str]) -> str:
    return subprocess.check_output(args, text=True)


def list_xlsx(prefix: str) -> list[str]:
    output = run_text(["gcloud", "storage", "ls", "--recursive", f"{prefix}**"])
    return sorted(line.strip() for line in output.splitlines() if line.endswith(".xlsx"))


def parse_cell_id(uri: str) -> int | None:
    match = CELL_PATTERN.search(uri)
    return int(match.group(1)) if match else None


def read_canonical_grid() -> list[dict[str, str]]:
    raw = run_text(["gcloud", "storage", "cat", CANONICAL_GRID_URI])
    return list(csv.DictReader(io.StringIO(raw)))


def choose_samples(cell_ids: list[int]) -> list[int]:
    ordered = sorted(cell_ids)
    if not ordered:
        return []
    last = len(ordered) - 1
    return sorted({ordered[round(q * last)] for q in SAMPLE_QUANTILES})


def profile_workbook(local_path: Path, surface: str, uri: str, cell_id: int) -> dict[str, Any]:
    workbook = load_workbook(local_path, read_only=True, data_only=True)
    if "Output" not in workbook.sheetnames:
        result = {
            "surface": surface,
            "cell_id": cell_id,
            "source_uri": uri,
            "sheet_names": workbook.sheetnames,
            "error": "missing Output sheet",
        }
        workbook.close()
        return result

    sheet = workbook["Output"]
    rows = list(sheet.iter_rows(values_only=True))
    headers = list(rows[0])
    data_rows = rows[1:]
    header_indexes = {str(header): index for index, header in enumerate(headers)}
    header_json = json.dumps(headers, separators=(",", ":"), ensure_ascii=False)
    dimensions = {}
    for column in (
        "ticcsSubClass",
        "ticcsSubClassName",
        "assetId",
        "countryCode",
        "climateZone",
        "geolocationCoordinates",
        "scenario",
        "timeHorizon",
        "hazard",
        "indicator",
        "indicatorUnit",
        "subrisk",
    ):
        if column not in header_indexes:
            continue
        index = header_indexes[column]
        values = []
        for row in data_rows:
            value = row[index]
            if value is not None and value not in values:
                values.append(value)
        dimensions[column] = sorted(values, key=str)

    null_counts = {
        str(header): sum(row[index] is None for row in data_rows)
        for index, header in enumerate(headers)
    }

    if "subrisk" in header_indexes:
        grain_columns = ["assetId", "scenario", "timeHorizon", "subrisk", "indicator"]
        repeated_group_columns = ["assetId", "scenario", "timeHorizon", "subrisk"]
        repeated_measure = "adjustedSubriskRevenueImpact"
    else:
        grain_columns = ["assetId", "scenario", "timeHorizon", "hazard", "indicator"]
        repeated_group_columns = ["assetId", "scenario", "timeHorizon"]
        repeated_measure = "adjustedTotalDamage"

    grain_indexes = [header_indexes[column] for column in grain_columns]
    row_keys = [tuple(row[index] for index in grain_indexes) for row in data_rows]
    repeated_groups: dict[tuple[Any, ...], set[Any]] = {}
    repeated_indexes = [header_indexes[column] for column in repeated_group_columns]
    measure_index = header_indexes[repeated_measure]
    for row in data_rows:
        group_key = tuple(row[index] for index in repeated_indexes)
        repeated_groups.setdefault(group_key, set()).add(row[measure_index])
    inconsistent_groups = sum(len(values) > 1 for values in repeated_groups.values())

    result = {
        "surface": surface,
        "cell_id": cell_id,
        "source_uri": uri,
        "sheet_names": workbook.sheetnames,
        "rows": len(data_rows),
        "columns": len(headers),
        "headers": headers,
        "schema_hash_sha256": hashlib.sha256(header_json.encode("utf-8")).hexdigest(),
        "dimensions": dimensions,
        "null_counts": null_counts,
        "source_grain_columns": grain_columns,
        "source_grain_unique_count": len(set(row_keys)),
        "source_grain_duplicate_count": len(row_keys) - len(set(row_keys)),
        "repeated_measure": repeated_measure,
        "repeated_measure_group_columns": repeated_group_columns,
        "repeated_measure_inconsistent_group_count": inconsistent_groups,
    }
    workbook.close()
    return result


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    canonical_rows = read_canonical_grid()
    canonical_by_id = {int(row["cell_id"]): row for row in canonical_rows}
    canonical_ids = set(canonical_by_id)

    surface_uris: dict[str, list[str]] = {}
    surface_ids: dict[str, set[int]] = {}
    inventory: dict[str, Any] = {}

    for name, prefix in SURFACES.items():
        uris = list_xlsx(prefix)
        parsed = [parse_cell_id(uri) for uri in uris]
        unreadable = [uri for uri, cell_id in zip(uris, parsed) if cell_id is None]
        valid_ids = [cell_id for cell_id in parsed if cell_id is not None]
        counts = Counter(valid_ids)
        duplicates = sorted(cell_id for cell_id, count in counts.items() if count > 1)
        cell_set = set(valid_ids)
        surface_uris[name] = uris
        surface_ids[name] = cell_set
        inventory[name] = {
            "prefix": prefix,
            "xlsx_file_count": len(uris),
            "unique_cell_count": len(cell_set),
            "unreadable_filename_count": len(unreadable),
            "duplicate_cell_count": len(duplicates),
            "missing_canonical_count": len(canonical_ids - cell_set),
            "extra_noncanonical_count": len(cell_set - canonical_ids),
            "missing_canonical_cells": sorted(canonical_ids - cell_set),
            "extra_noncanonical_cells": sorted(cell_set - canonical_ids),
        }

    comparisons = {}
    names = sorted(surface_ids)
    for index, left in enumerate(names):
        for right in names[index + 1 :]:
            comparisons[f"{left}__vs__{right}"] = {
                "only_left_count": len(surface_ids[left] - surface_ids[right]),
                "only_right_count": len(surface_ids[right] - surface_ids[left]),
            }

    summary = {
        "canonical_grid_uri": CANONICAL_GRID_URI,
        "canonical_row_count": len(canonical_rows),
        "canonical_unique_cell_count": len(canonical_ids),
        "surfaces": inventory,
        "pairwise_comparisons": comparisons,
    }
    (OUTPUT_DIR / "inventory_summary.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )

    solar_missing = sorted(
        canonical_ids
        - (surface_ids["solar_physical"] & surface_ids["solar_transition"])
    )
    with (OUTPUT_DIR / "solar_missing_cells.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(
            handle, fieldnames=list(canonical_rows[0]), lineterminator="\n"
        )
        writer.writeheader()
        for cell_id in solar_missing:
            writer.writerow(canonical_by_id[cell_id])

    profiles = []
    with tempfile.TemporaryDirectory(prefix="scr-four-surface-") as temp_dir:
        temp_root = Path(temp_dir)
        for surface, uris in surface_uris.items():
            uri_by_cell = {
                parse_cell_id(uri): uri for uri in uris if parse_cell_id(uri) is not None
            }
            for cell_id in choose_samples(sorted(surface_ids[surface])):
                uri = uri_by_cell[cell_id]
                local_path = temp_root / f"{surface}_{cell_id}.xlsx"
                subprocess.check_call(["gcloud", "storage", "cp", uri, str(local_path)])
                profiles.append(profile_workbook(local_path, surface, uri, cell_id))

    (OUTPUT_DIR / "schema_profile.json").write_text(
        json.dumps(profiles, indent=2, ensure_ascii=False, default=str) + "\n",
        encoding="utf-8",
    )

    with (OUTPUT_DIR / "schema_sample_summary.csv").open(
        "w", newline="", encoding="utf-8"
    ) as handle:
        fields = [
            "surface",
            "cell_id",
            "rows",
            "columns",
            "schema_hash_sha256",
            "scenarios",
            "horizons",
            "hazards",
            "indicators",
            "subrisks",
            "source_uri",
            "error",
        ]
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        for profile in profiles:
            dimensions = profile.get("dimensions", {})
            writer.writerow(
                {
                    "surface": profile["surface"],
                    "cell_id": profile["cell_id"],
                    "rows": profile.get("rows"),
                    "columns": profile.get("columns"),
                    "schema_hash_sha256": profile.get("schema_hash_sha256"),
                    "scenarios": json.dumps(dimensions.get("scenario", []), default=str),
                    "horizons": json.dumps(dimensions.get("timeHorizon", []), default=str),
                    "hazards": json.dumps(dimensions.get("hazard", []), default=str),
                    "indicators": json.dumps(dimensions.get("indicator", []), default=str),
                    "subrisks": json.dumps(dimensions.get("subrisk", []), default=str),
                    "source_uri": profile["source_uri"],
                    "error": profile.get("error"),
                }
            )

    print(json.dumps(summary, indent=2, ensure_ascii=False))
    print(f"Wrote evidence to {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
