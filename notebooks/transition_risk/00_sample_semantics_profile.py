#!/usr/bin/env python3
"""Profile current SCR transition workbooks without assuming financial semantics."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import subprocess
import tempfile
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


SURFACES = ("solar_transition", "wind_transition")
CELL_ID_RE = re.compile(r"Cell_(\d+)_")
EXPECTED_INDICATORS = (
    "Carbon price",
    "Revenue growth",
    "Scope 1&2 emissions intensity",
    "Scope 3 emissions intensity",
)
EXPECTED_SUBRISKS = ("Direct Carbon Cost", "Market Demand Shifts")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--schema-profile", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    return parser.parse_args()


def normalize(value):
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


def stable_hash(values: list[str]) -> str:
    return hashlib.sha256("\n".join(values).encode("utf-8")).hexdigest()


def read_workbook(path: Path) -> tuple[dict, list[dict]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    readme = workbook["ReadMe"]
    definitions = {}
    asset_metadata = {}
    for row in readme.iter_rows(values_only=True):
        values = list(row)
        populated = [value for value in values if value not in (None, "")]
        if len(populated) >= 2 and populated[0] in {
            "assetId",
            "assetName",
            "assetType",
            "country",
            "assetValue",
            "assetRevenue",
        }:
            asset_metadata[populated[0]] = populated[1]
        if len(values) < 6 or values[1] in (None, "Header Name"):
            continue
        header = values[1]
        if isinstance(header, str) and header in {
            "scenario",
            "timeHorizon",
            "indicatorValue",
            "adjustedIndicatorValue",
            "subriskRevenueImpact",
            "adjustedSubriskRevenueImpact",
            "subriskExposureRating",
            "adjustedSubriskExposureRating",
            "transitionExposureRating",
            "adjustedTransitionExposureRating",
        }:
            definitions[header] = {
                "definition": normalize(values[2]),
                "type": normalize(values[3]),
                "range": normalize(values[4]),
                "unit": normalize(values[5]),
            }

    output = workbook["Output"]
    row_iter = output.iter_rows(values_only=True)
    headers = [str(value) for value in next(row_iter)]
    rows = [
        {header: normalize(value) for header, value in zip(headers, values)}
        for values in row_iter
        if any(value is not None for value in values)
    ]
    workbook.close()
    return {
        "definitions": definitions,
        "headers": headers,
        "asset_metadata": asset_metadata,
    }, rows


def unique(values) -> list:
    return sorted(set(values), key=lambda item: (str(type(item)), str(item)))


def main() -> None:
    args = parse_args()
    profiles = json.loads(args.schema_profile.read_text(encoding="utf-8"))
    solar_sample = [item for item in profiles if item["surface"] == "solar_transition"]
    if len(solar_sample) != 7:
        raise ValueError("Expected seven deterministic Solar transition profiles")
    wind_profile = next(
        item for item in profiles if item["surface"] == "wind_transition"
    )
    wind_prefix = wind_profile["source_uri"].rsplit("/", 1)[0] + "/"
    listing = subprocess.check_output(
        ["gcloud", "storage", "ls", "--recursive", f"{wind_prefix}**"],
        text=True,
    )
    wind_uri_by_cell = {}
    for uri in listing.splitlines():
        match = CELL_ID_RE.search(uri)
        if match and uri.endswith(".xlsx"):
            wind_uri_by_cell[int(match.group(1))] = uri
    sample_cell_ids = [item["cell_id"] for item in solar_sample]
    selected = {
        "solar_transition": solar_sample,
        "wind_transition": [
            {"surface": "wind_transition", "cell_id": cell_id, "source_uri": wind_uri_by_cell[cell_id]}
            for cell_id in sample_cell_ids
        ],
    }

    parsed: dict[str, list[dict]] = {surface: [] for surface in SURFACES}
    readme_by_surface = {}
    source_uris = []
    with tempfile.TemporaryDirectory(prefix="scr-transition-profile-") as temp_dir:
        temp_root = Path(temp_dir)
        for surface, items in selected.items():
            for item in items:
                uri = item["source_uri"]
                source_uris.append(uri)
                local = temp_root / f"{surface}_{item['cell_id']}.xlsx"
                subprocess.check_call(
                    ["gcloud", "--quiet", "storage", "cp", uri, str(local)],
                    stdout=subprocess.DEVNULL,
                )
                metadata, rows = read_workbook(local)
                if len(rows) != 216 or len(metadata["headers"]) != 20:
                    raise ValueError(f"Unexpected transition shape for {uri}")
                readme_by_surface.setdefault(surface, metadata["definitions"])
                for row in rows:
                    row["surface"] = surface
                    row["cell_id"] = int(item["cell_id"])
                    row["source_uri"] = uri
                    row["readme_asset_value"] = metadata["asset_metadata"].get("assetValue")
                    row["readme_asset_revenue"] = metadata["asset_metadata"].get("assetRevenue")
                parsed[surface].extend(rows)

    report = {
        "sample_design": {
            "workbooks_per_surface": 7,
            "total_workbooks": 14,
            "source_uri_list_sha256": stable_hash(sorted(source_uris)),
            "cell_ids_by_surface": {
                surface: [item["cell_id"] for item in items]
                for surface, items in selected.items()
            },
        },
        "readme_contract": readme_by_surface,
        "surfaces": {},
        "cross_asset": {},
        "interpretation_findings": [],
    }
    compact_rows = []

    for surface, rows in parsed.items():
        cell_ids = sorted({row["cell_id"] for row in rows})
        scenarios = unique(row["scenario"] for row in rows)
        horizons = unique(row["timeHorizon"] for row in rows)
        indicators = unique(row["indicator"] for row in rows)
        subrisks = unique(row["subrisk"] for row in rows)
        grain = Counter(
            (
                row["cell_id"],
                row["assetId"],
                row["scenario"],
                row["timeHorizon"],
                row["subrisk"],
                row["indicator"],
            )
            for row in rows
        )

        repeated_consistency = {}
        for measure in (
            "subriskRevenueImpact",
            "adjustedSubriskRevenueImpact",
            "subriskExposureRating",
            "adjustedSubriskExposureRating",
            "transitionExposureRating",
            "adjustedTransitionExposureRating",
        ):
            groups = defaultdict(set)
            group_columns = (
                "cell_id",
                "scenario",
                "timeHorizon",
                "subrisk",
            ) if measure.startswith(("subrisk", "adjustedSubrisk")) else (
                "cell_id",
                "scenario",
                "timeHorizon",
            )
            for row in rows:
                groups[tuple(row[column] for column in group_columns)].add(row[measure])
            repeated_consistency[measure] = {
                "group_count": len(groups),
                "inconsistent_group_count": sum(len(values) > 1 for values in groups.values()),
            }

        impacts = [
            float(row["adjustedSubriskRevenueImpact"])
            for row in rows
            if row["adjustedSubriskRevenueImpact"] is not None
        ]
        impact_pairs = [
            (row["subriskRevenueImpact"], row["adjustedSubriskRevenueImpact"])
            for row in rows
        ]
        rating_pairs = [
            (row["transitionExposureRating"], row["adjustedTransitionExposureRating"])
            for row in rows
        ]

        location_variation = {}
        for value_field, group_columns in (
            ("indicatorValue", ("scenario", "timeHorizon", "indicator")),
            (
                "adjustedSubriskRevenueImpact",
                ("scenario", "timeHorizon", "subrisk"),
            ),
            (
                "adjustedSubriskExposureRating",
                ("scenario", "timeHorizon", "subrisk"),
            ),
            (
                "adjustedTransitionExposureRating",
                ("scenario", "timeHorizon"),
            ),
        ):
            groups = defaultdict(dict)
            for row in rows:
                key = tuple(row[column] for column in group_columns)
                groups[key][row["cell_id"]] = row[value_field]
            complete = [values for values in groups.values() if len(values) == len(cell_ids)]
            varying = [
                (key, values)
                for key, values in groups.items()
                if len(values) == len(cell_ids) and len(set(values.values())) > 1
            ]
            location_variation[value_field] = {
                "complete_comparison_groups": len(complete),
                "groups_varying_across_sample_cells": len(varying),
                "examples": [
                    {
                        "group": list(key),
                        "values_by_cell": {
                            str(cell_id): value
                            for cell_id, value in sorted(values.items())
                        },
                    }
                    for key, values in varying[:5]
                ],
            }

        # One compact row per cell/scenario/horizon. Repeated subrisk measures are
        # deduplicated and indicators are pivoted into explicit columns.
        compact = {}
        for row in rows:
            key = (row["cell_id"], row["scenario"], row["timeHorizon"])
            candidate = compact.setdefault(
                key,
                {
                    "surface": surface,
                    "cell_id": row["cell_id"],
                    "asset_id": row["assetId"],
                    "asset_type": row["ticcsSubClassName"],
                    "ticcs_subclass": row["ticcsSubClass"],
                    "report_date": row["reportDate"].date().isoformat(),
                    "readme_asset_value": row["readme_asset_value"],
                    "readme_asset_revenue": row["readme_asset_revenue"],
                    "scenario": row["scenario"],
                    "horizon": int(row["timeHorizon"]),
                    "source_uri": row["source_uri"],
                },
            )
            indicator_key = {
                "Carbon price": "carbon_price",
                "Revenue growth": "revenue_growth",
                "Scope 1&2 emissions intensity": "scope_1_2_intensity",
                "Scope 3 emissions intensity": "scope_3_intensity",
            }[row["indicator"]]
            candidate[indicator_key] = row["indicatorValue"]
            candidate[f"{indicator_key}_unit"] = row["indicatorUnit"]
            subrisk_key = {
                "Direct Carbon Cost": "direct_carbon_cost",
                "Market Demand Shifts": "market_demand_shifts",
            }[row["subrisk"]]
            for name, value in (
                (f"{subrisk_key}_impact", row["adjustedSubriskRevenueImpact"]),
                (f"{subrisk_key}_rating", row["adjustedSubriskExposureRating"]),
                ("overall_transition_rating", row["adjustedTransitionExposureRating"]),
            ):
                if name in candidate and candidate[name] != value:
                    raise ValueError(f"Inconsistent repeated compact value: {key}/{name}")
                candidate[name] = value
        compact_rows.extend(compact.values())

        report["surfaces"][surface] = {
            "sample_workbook_count": len(cell_ids),
            "source_row_count": len(rows),
            "source_rows_per_workbook": len(rows) // len(cell_ids),
            "source_grain_duplicate_count": sum(count - 1 for count in grain.values()),
            "scenarios": scenarios,
            "horizons": horizons,
            "indicators": indicators,
            "subrisks": subrisks,
            "report_dates_by_cell": {
                str(cell_id): sorted({
                    row["reportDate"].date().isoformat()
                    for row in rows
                    if row["cell_id"] == cell_id
                })
                for cell_id in cell_ids
            },
            "readme_asset_values": unique(row["readme_asset_value"] for row in rows),
            "readme_asset_revenues": unique(row["readme_asset_revenue"] for row in rows),
            "compact_row_count": len(compact),
            "compact_rows_per_workbook": len(compact) // len(cell_ids),
            "impact_distribution": {
                "minimum": min(impacts),
                "maximum": max(impacts),
                "negative_row_count": sum(value < 0 for value in impacts),
                "above_one_abs_row_count": sum(abs(value) > 1 for value in impacts),
                "above_100_abs_row_count": sum(abs(value) > 100 for value in impacts),
            },
            "raw_equals_adjusted_impact_share": sum(a == b for a, b in impact_pairs)
            / len(impact_pairs),
            "raw_equals_adjusted_overall_rating_share": sum(a == b for a, b in rating_pairs)
            / len(rating_pairs),
            "repeated_measure_consistency": repeated_consistency,
            "location_variation": location_variation,
            "checks": {
                "shape_216_by_20": len(rows) == len(cell_ids) * 216,
                "source_grain_unique": all(count == 1 for count in grain.values()),
                "nine_scenarios": len(scenarios) == 9,
                "six_horizons": len(horizons) == 6,
                "expected_indicators": tuple(indicators) == tuple(sorted(EXPECTED_INDICATORS)),
                "expected_subrisks": tuple(subrisks) == tuple(sorted(EXPECTED_SUBRISKS)),
                "compact_grain_54_per_workbook": len(compact) == len(cell_ids) * 54,
                "repeated_measures_consistent": all(
                    item["inconsistent_group_count"] == 0
                    for item in repeated_consistency.values()
                ),
            },
        }

    solar = {
        (row["cell_id"], row["scenario"], row["horizon"]): row
        for row in compact_rows
        if row["surface"] == "solar_transition"
    }
    wind = {
        (row["cell_id"], row["scenario"], row["horizon"]): row
        for row in compact_rows
        if row["surface"] == "wind_transition"
    }
    common_keys = sorted(set(solar) & set(wind))
    comparisons = {}
    for field in (
        "carbon_price",
        "revenue_growth",
        "scope_1_2_intensity",
        "scope_3_intensity",
        "direct_carbon_cost_impact",
        "market_demand_shifts_impact",
        "overall_transition_rating",
    ):
        comparisons[field] = {
            "comparable_rows": len(common_keys),
            "equal_rows": sum(solar[key].get(field) == wind[key].get(field) for key in common_keys),
        }
    report["cross_asset"] = comparisons

    readme = readme_by_surface["solar_transition"]
    report["interpretation_findings"] = [
        {
            "finding": "SCR labels adjustedSubriskRevenueImpact as an annualized average impact on asset revenue / OpEx with unit percent.",
            "evidence": readme["adjustedSubriskRevenueImpact"],
        },
        {
            "finding": "The documented 0-1 range is inconsistent with current output values, including magnitudes above 1; scale must be validated before cash-flow application.",
            "evidence": {
                surface: report["surfaces"][surface]["impact_distribution"]
                for surface in SURFACES
            },
        },
        {
            "finding": "The ReadMe scenario and horizon descriptions are stale relative to the exported Output sheet.",
            "evidence": {
                "readme_scenario_range": readme["scenario"]["range"],
                "readme_horizon_range": readme["timeHorizon"]["range"],
                "observed_scenarios": report["surfaces"]["solar_transition"]["scenarios"],
                "observed_horizons": report["surfaces"]["solar_transition"]["horizons"],
            },
        },
        {
            "finding": "The ReadMe defines adjustedIndicatorValue, but that field is absent from the current 20-column Output sheet.",
            "evidence": {
                "definition": readme["adjustedIndicatorValue"],
                "output_headers": selected["solar_transition"][0]["headers"],
            },
        },
    ]

    if not all(
        all(surface["checks"].values()) for surface in report["surfaces"].values()
    ):
        raise ValueError("Transition sample profile checks failed")

    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / "transition_sample_semantics_profile.json").write_text(
        json.dumps(report, indent=2, default=str) + "\n", encoding="utf-8"
    )
    fieldnames = sorted({field for row in compact_rows for field in row})
    with (args.output_dir / "transition_compact_sample.csv").open(
        "w", newline="", encoding="utf-8"
    ) as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, lineterminator="\n")
        writer.writeheader()
        writer.writerows(sorted(compact_rows, key=lambda row: (
            row["surface"], row["cell_id"], row["scenario"], row["horizon"]
        )))
    print(json.dumps({
        "status": "passed",
        "workbooks": 14,
        "source_rows": sum(len(rows) for rows in parsed.values()),
        "compact_rows": len(compact_rows),
        "solar_impact_range": report["surfaces"]["solar_transition"]["impact_distribution"],
        "wind_impact_range": report["surfaces"]["wind_transition"]["impact_distribution"],
    }))


if __name__ == "__main__":
    main()
