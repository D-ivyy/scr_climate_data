#!/usr/bin/env python3
"""Detect content regimes in deterministic Solar/Wind transition samples."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import subprocess
import tempfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


CELL_ID_RE = re.compile(r"Cell_(\d+)_")
SURFACES = {
    "solar_transition": "gs://infrasure-scr-data/Solar-asset/transition_risk/",
    "wind_transition": "gs://infrasure-scr-data/onshore_wind/transition_risk/",
}
FINGERPRINT_COLUMNS = (
    "scenario",
    "timeHorizon",
    "indicator",
    "indicatorUnit",
    "indicatorValue",
    "subrisk",
    "subriskRevenueImpact",
    "adjustedSubriskRevenueImpact",
    "subriskExposureRating",
    "adjustedSubriskExposureRating",
    "transitionExposureRating",
    "adjustedTransitionExposureRating",
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--sample-size", type=int, default=128)
    return parser.parse_args()


def choose_evenly(values: list[str], size: int) -> list[str]:
    if size >= len(values):
        return values
    indexes = sorted({round(index * (len(values) - 1) / (size - 1)) for index in range(size)})
    return [values[index] for index in indexes]


def sha256_json(value) -> str:
    payload = json.dumps(value, sort_keys=True, separators=(",", ":"), default=str)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def parse_workbook(path: Path) -> dict:
    match = CELL_ID_RE.search(path.name)
    if not match:
        raise ValueError(f"Cannot parse cell ID from {path.name}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Output"]
    row_iter = sheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(row_iter)]
    rows = [dict(zip(headers, values)) for values in row_iter]
    workbook.close()
    content = [
        [row[column] for column in FINGERPRINT_COLUMNS]
        for row in sorted(
            rows,
            key=lambda row: (
                row["scenario"],
                int(row["timeHorizon"]),
                row["subrisk"],
                row["indicator"],
            ),
        )
    ]
    sentinel = {}
    for row in rows:
        if row["scenario"] == "Expected" and str(row["timeHorizon"]) == "2030":
            if row["indicator"] == "Revenue growth":
                sentinel["expected_2030_revenue_growth"] = row["indicatorValue"]
            if row["subrisk"] == "Market Demand Shifts":
                sentinel["expected_2030_market_demand_impact"] = row[
                    "adjustedSubriskRevenueImpact"
                ]
                sentinel["expected_2030_market_demand_rating"] = row[
                    "adjustedSubriskExposureRating"
                ]
        if row["scenario"] == "Expected" and str(row["timeHorizon"]) == "2025":
            sentinel["expected_2025_overall_rating"] = row[
                "adjustedTransitionExposureRating"
            ]
    first = rows[0]
    return {
        "cell_id": int(match.group(1)),
        "asset_id": first["assetId"],
        "report_date": first["reportDate"].date().isoformat(),
        "ticcs_subclass": first["ticcsSubClass"],
        "economic_content_fingerprint": sha256_json(content),
        **sentinel,
    }


def main() -> None:
    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    results = []
    inventory_counts = {}
    sample_uri_hashes = {}
    with tempfile.TemporaryDirectory(prefix="scr-transition-fingerprint-") as temp_dir:
        temp_root = Path(temp_dir)
        for surface, prefix in SURFACES.items():
            listing = subprocess.check_output(
                ["gcloud", "storage", "ls", "--recursive", f"{prefix}**"], text=True
            )
            uris = sorted(line for line in listing.splitlines() if line.endswith(".xlsx"))
            inventory_counts[surface] = len(uris)
            chosen = choose_evenly(uris, args.sample_size)
            sample_uri_hashes[surface] = sha256_json(chosen)
            local_dir = temp_root / surface
            local_dir.mkdir()
            for offset in range(0, len(chosen), 64):
                subprocess.check_call(
                    [
                        "gcloud",
                        "--quiet",
                        "storage",
                        "cp",
                        *chosen[offset : offset + 64],
                        str(local_dir) + "/",
                    ],
                    stdout=subprocess.DEVNULL,
                )
            for uri in chosen:
                parsed = parse_workbook(local_dir / Path(uri).name)
                parsed.update({"surface": surface, "source_uri": uri})
                results.append(parsed)

    report = {
        "sample_size_requested_per_surface": args.sample_size,
        "inventory_counts": inventory_counts,
        "sample_uri_list_sha256": sample_uri_hashes,
        "surfaces": {},
        "cross_surface_shared_fingerprint_count": 0,
    }
    fingerprints_by_surface = {}
    for surface in SURFACES:
        rows = sorted(
            [row for row in results if row["surface"] == surface],
            key=lambda row: row["cell_id"],
        )
        grouped = defaultdict(list)
        for row in rows:
            grouped[row["economic_content_fingerprint"]].append(row)
        fingerprints_by_surface[surface] = set(grouped)
        transitions = sum(
            left["economic_content_fingerprint"]
            != right["economic_content_fingerprint"]
            for left, right in zip(rows, rows[1:])
        )
        report["surfaces"][surface] = {
            "sample_count": len(rows),
            "distinct_economic_content_fingerprints": len(grouped),
            "economic_fingerprint_transitions_in_cell_order": transitions,
            "report_dates": sorted({row["report_date"] for row in rows}),
            "ticcs_subclasses": sorted({row["ticcs_subclass"] for row in rows}),
            "fingerprints": [
                {
                    "economic_content_fingerprint": fingerprint,
                    "sample_count": len(items),
                    "minimum_cell_id": min(row["cell_id"] for row in items),
                    "maximum_cell_id": max(row["cell_id"] for row in items),
                    "sample_cell_ids": [row["cell_id"] for row in items[:12]],
                    "asset_id_range": [
                        min(row["asset_id"] for row in items),
                        max(row["asset_id"] for row in items),
                    ],
                    "sentinel": {
                        key: items[0].get(key)
                        for key in (
                            "expected_2030_revenue_growth",
                            "expected_2030_market_demand_impact",
                            "expected_2030_market_demand_rating",
                            "expected_2025_overall_rating",
                        )
                    },
                }
                for fingerprint, items in sorted(
                    grouped.items(), key=lambda item: (-len(item[1]), item[0])
                )
            ],
        }
    report["cross_surface_shared_fingerprint_count"] = len(
        fingerprints_by_surface["solar_transition"]
        & fingerprints_by_surface["wind_transition"]
    )

    (args.output_dir / "transition_template_fingerprint_sample.json").write_text(
        json.dumps(report, indent=2) + "\n", encoding="utf-8"
    )
    with (args.output_dir / "transition_template_fingerprint_sample.csv").open(
        "w", newline="", encoding="utf-8"
    ) as handle:
        fieldnames = list(results[0])
        writer = csv.DictWriter(handle, fieldnames=fieldnames, lineterminator="\n")
        writer.writeheader()
        writer.writerows(sorted(results, key=lambda row: (row["surface"], row["cell_id"])))
    print(json.dumps({
        "status": "passed",
        "sample_counts": {
            surface: report["surfaces"][surface]["sample_count"] for surface in SURFACES
        },
        "distinct_fingerprints": {
            surface: report["surfaces"][surface][
                "distinct_economic_content_fingerprints"
            ]
            for surface in SURFACES
        },
        "shared_fingerprints": report["cross_surface_shared_fingerprint_count"],
    }))


if __name__ == "__main__":
    main()
