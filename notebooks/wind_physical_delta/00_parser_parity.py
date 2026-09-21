#!/usr/bin/env python3
"""Validate the production OOXML extractor against openpyxl for Wind samples."""

from __future__ import annotations

import importlib.util
import json
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
PROFILE_PATH = (
    REPO_ROOT
    / "notebooks/scr_four_surface_inventory/outputs/schema_profile.json"
)
EXTRACTOR_PATH = REPO_ROOT / "cloudrun/scr_physical_delta/run_task.py"
OUTPUT_PATH = Path(__file__).resolve().parent / "outputs/parser_parity.json"


def load_extractor():
    spec = importlib.util.spec_from_file_location("scr_physical_delta_run_task", EXTRACTOR_PATH)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Unable to load extractor from {EXTRACTOR_PATH}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def distinct(values: list[Any]) -> Any:
    present = []
    for value in values:
        if value not in (None, "") and value not in present:
            present.append(value)
    if len(present) > 1:
        raise ValueError(f"Expected one repeated adjustedTotalDamage value, found {present}")
    return present[0] if present else None


def openpyxl_reference(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Output" not in workbook.sheetnames:
        raise ValueError("Missing Output worksheet")
    rows = list(workbook["Output"].iter_rows(values_only=True))
    headers = list(rows[0])
    indexes = {str(header): index for index, header in enumerate(headers)}
    required = {
        "assetId",
        "ticcsSubClass",
        "ticcsSubClassName",
        "scenario",
        "timeHorizon",
        "adjustedTotalDamage",
    }
    missing = sorted(required - indexes.keys())
    if missing:
        raise ValueError(f"Missing required headers: {missing}")

    grouped: dict[str, dict[str, list[Any]]] = {}
    for row in rows[1:]:
        scenario = row[indexes["scenario"]]
        horizon = row[indexes["timeHorizon"]]
        if scenario in (None, "") or horizon in (None, ""):
            continue
        grouped.setdefault(str(scenario), {}).setdefault(str(horizon), []).append(
            row[indexes["adjustedTotalDamage"]]
        )
    totals = {
        scenario: {horizon: distinct(values) for horizon, values in horizons.items()}
        for scenario, horizons in grouped.items()
    }
    identity = {
        column: sorted(
            {
                row[indexes[column]]
                for row in rows[1:]
                if row[indexes[column]] not in (None, "")
            },
            key=str,
        )
        for column in ("assetId", "ticcsSubClass", "ticcsSubClassName")
    }
    workbook.close()
    return {
        "rowCount": len(rows) - 1,
        "headers": headers,
        "totals": totals,
        "identity": identity,
    }


def main() -> None:
    profiles = json.loads(PROFILE_PATH.read_text(encoding="utf-8"))
    samples = sorted(
        (profile for profile in profiles if profile["surface"] == "wind_physical"),
        key=lambda profile: profile["cell_id"],
    )
    if not samples:
        raise ValueError("No Wind physical samples found; run the four-surface inventory first")

    extractor = load_extractor()
    results = []
    with tempfile.TemporaryDirectory(prefix="scr-wind-parser-parity-") as temp_dir:
        root = Path(temp_dir)
        for sample in samples:
            uri = sample["source_uri"]
            local_path = root / Path(uri).name
            subprocess.check_call(["gcloud", "storage", "cp", uri, str(local_path)])
            production = extractor.extract_workbook(local_path)
            reference = openpyxl_reference(local_path)
            checks = {
                "production_status_ok": production.get("status") == "ok",
                "cell_id_matches": production.get("cellId") == sample["cell_id"],
                "row_count_matches": production.get("rowCount") == reference["rowCount"],
                "headers_match": production.get("headers") == reference["headers"],
                "totals_match": production.get("totals") == reference["totals"],
                "wind_ticcs_matches": reference["identity"]["ticcsSubClass"] == ["IC701010"],
            }
            results.append(
                {
                    "cell_id": sample["cell_id"],
                    "source_uri": uri,
                    "identity": reference["identity"],
                    "checks": checks,
                    "status": "passed" if all(checks.values()) else "failed",
                }
            )

    summary = {
        "status": "passed" if all(item["status"] == "passed" for item in results) else "failed",
        "sample_count": len(results),
        "passed_count": sum(item["status"] == "passed" for item in results),
        "failed_count": sum(item["status"] != "passed" for item in results),
        "extractor": str(EXTRACTOR_PATH.relative_to(REPO_ROOT)),
        "results": results,
    }
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({key: summary[key] for key in ("status", "sample_count", "passed_count", "failed_count")}))
    if summary["status"] != "passed":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
